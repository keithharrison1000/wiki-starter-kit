"""Convert a spreadsheet (.xlsx) to markdown: one table per worksheet.

Deliberately simple relative to pdf2md — no layout inference, no attempt to
split a sheet into multiple logical tables. Each worksheet becomes exactly
one markdown table, trimmed to its actual used range (openpyxl's own
max_row/max_column is unreliable — formatting-only cells routinely inflate
it far past where real content ends), with the sheet's own first row used
as the table header. A sheet that mixes a summary block, a data table, and
a legend under a handful of blank rows (a common human-authored layout)
still renders faithfully as one table; blank rows just become empty rows in
the output rather than visual section breaks. Good enough for a Source
concept's job — discoverable and citable — not a final presentation.
"""

from __future__ import annotations

import datetime
from pathlib import Path

import openpyxl


def _used_bounds(ws) -> tuple[int, int]:
    """The last row and column that actually hold a value, scanning the
    whole sheet rather than trusting ws.max_row/max_column."""
    max_row = 0
    max_col = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                max_row = max(max_row, cell.row)
                max_col = max(max_col, cell.column)
    return max_row, max_col


def _is_percent_format(number_format: str) -> bool:
    return "%" in number_format


def _format_cell(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        value = value.date()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        if _is_percent_format(cell.number_format or ""):
            pct = value * 100
            text = f"{pct:.1f}".rstrip("0").rstrip(".")
            return f"{text}%"
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        if isinstance(value, int):
            return f"{value:,}"
        return f"{value:,.2f}"
    text = str(value).strip()
    return text.replace("|", "\\|").replace("\n", "<br>")


def convert(input_path: str | Path) -> str:
    """Convert every worksheet in an .xlsx workbook to a markdown table,
    concatenated under one heading per sheet."""
    path = Path(input_path)
    wb = openpyxl.load_workbook(path, data_only=True)

    sections: list[str] = []
    for ws in wb.worksheets:
        max_row, max_col = _used_bounds(ws)
        if max_row == 0 or max_col == 0:
            continue  # a genuinely empty sheet

        rows = [
            [_format_cell(ws.cell(row=r, column=c)) for c in range(1, max_col + 1)]
            for r in range(1, max_row + 1)
        ]

        header, body = rows[0], rows[1:]
        lines = [f"## {ws.title}", ""]
        lines.append("| " + " | ".join(header) + " |")
        lines.append("|" + "|".join(["---"] * max_col) + "|")
        for row in body:
            lines.append("| " + " | ".join(row) + " |")
        sections.append("\n".join(lines))

    return "\n\n".join(sections) + "\n"
